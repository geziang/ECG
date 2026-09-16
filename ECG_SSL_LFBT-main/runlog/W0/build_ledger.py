# -*- coding: utf-8 -*-
"""W0-5: 重建 LFBT实验方向清单.xlsx 台账(三 sheet),内容与 01/04 号 md 一致。"""
import sys, os
XLSX_SKILL_DIR = r"C:\Users\admin\.zcode\cli\plugins\cache\zcode-plugins-official\document-skills\0.1.4\skills\xlsx"
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")]:
    sys.path.insert(0, sub)
import base as B
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

B.use_palette_explicit("professional")

wb = Workbook()

# ---------- Sheet 1: 方向清单 ----------
ws = wb.active
ws.title = "方向清单"
B.setup_sheet(ws, title="SD-LFBT 实验方向清单(D1-D10 模块轴 + N1-N8 新轴)", last_col=7)
headers = ["编号", "轴/插槽", "一句话内容", "决策", "依据/联动", "来源文档"]
for c, h in enumerate(headers, 2):
    ws.cell(row=4, column=c, value=h)
B.style_header_row(ws, 4, 2, 7)
rows = [
    ["D1L", "导联轴(关系目标)", "inter-loss 目标矩阵改生理拓扑设定(II-III Einthoven 夹角、相邻胸导非零),只改目标值零新参数", "🟡 与 D3 同插槽竞争(S2)", "LGA 拉近式显著负 vs S1 对齐式弱正→坚持结构化对齐;需 NEG 负对照", "01 §4.3"],
    ["D1(重GAT)", "导联轴", "导联注意力/混合/token 交互", "❌ 弃", "E002/E003 双重否决", "01 §4.7"],
    ["D2", "导联轴", "偏相关", "⬇ 降级为 D3 ablation", "batch 128 协方差求逆噪声大", "01 §4.7"],
    ["D3", "导联轴(关系层)", "共源-特异解耦 z=[c;s], c 跨导联一致、s 去相关、c⊥s", "🟢 S2 主创新", "E007 证明两极端皆错,D3 为中间态;塌缩退 D1L", "01 §4.2"],
    ["D4", "时间轴(编码器)", "CNN 保形态 + Mamba 长程节律主干", "🟡 S3,带硬前提 G-D4 参数对齐对照", "E004 轻量适配器已终止,须换真主干", "01 §4.4"],
    ["D5", "时间轴", "时-频双分支", "⏸ 第二梯队", "E006-P1 前哨:AR-B3 弱正(FT100 +0.0027 3/3)、B0 无增益", "01 §4.7"],
    ["D6", "时间轴", "多尺度 multi-crop", "❌ 弃", "短窗正样本语义不一致+E003 同族不乐观", "01 §4.7"],
    ["D7", "目标轴(视图)", "时间掩码重建支路 L=L_BT+η·L_rec(mask0.5,η0.1)", "🟢 S4 冲点", "E005 证伪蒸馏式,重建式不同族;写作与 CoRe-ECG 划界", "01 §4.5"],
    ["D8", "目标轴(视图)", "物理噪声库+0.5-40Hz 频谱一致;评估改道 SNR 曲线+37条件错位+跨库", "🟢 S5 防御", "SNPH 不可得;判据宽松(回落≤0.3pt)", "01 §4.6"],
    ["D9", "目标轴(损失)", "VICReg 三项化(invariance+variance hinge+covariance)+白化,BN→LN", "🟢 S1 第一优先底座", "唯一正面解决 batch128 弱区间;胜条件=不劣于 BT 且 batch 曲线更平", "01 §4.1"],
    ["D10", "目标轴(优化)", "EMA 自蒸馏+原型", "❌ 弃", "E005 证伪师生蒸馏同族", "01 §4.7"],
    ["N1", "数据轴", "多库联合预训练(PTB-XL+CPSC+Georgia+PTB+INCART ≈3.87万,零下载)", "✅ 底座级语料轨并行(S0.5)", "ptb-xl 禁入预训练(与下游同源)", "01 §4.8"],
    ["N2", "数据轴", "参数可控仿真 ECG 混入预训练", "✅ S4 同期粗筛", "SimECG 证据", "01 §4.8"],
    ["N3", "采样轴", "同患者不同记录互为附加正对", "✅ D9 底座上择机插入", "低成本增广", "01 §4.8"],
    ["N4", "优化轴", "末段权重 EMA/soup", "✅ 免费附件默认开", "不改损失不加分支", "01 §4.8"],
    ["N5", "优化轴", "SAM/F-SAM 锐度感知", "⚠️ 第三梯队", "训练×2 成本,需防虚报", "01 §4.8"],
    ["N6", "表征深度", "projector 中间层辅助 BT 损失", "⚠️ D9 延伸项", "—", "01 §4.8"],
    ["N7", "鲁棒轴", "纯增强式导联置零一致性", "✅ 并入 D8", "E002 LeadMask 线索未关闭", "01 §4.8"],
    ["N8", "元数据轴", "age/sex 回归辅助头", "⚠️ 附录/防御", "牺牲纯 SSL 定位", "01 §4.8"],
]
for i, r in enumerate(rows):
    rn = 5 + i
    for c, v in enumerate(r, 2):
        ws.cell(row=rn, column=c, value=v)
    B.style_data_row(ws, rn, 2, 7, i)
B.auto_fit_columns(ws, header_row=4, data_start_row=5)
B.auto_fit_row_heights(ws, header_row=4, data_start_row=5, data_end_row=4 + len(rows))
ws.freeze_panes = "C5"

# ---------- Sheet 2: RunLog ----------
ws2 = wb.create_sheet("RunLog")
B.setup_sheet(ws2, title="实验 RunLog(E001-E008+,顺序执行,不选择性报告)", last_col=7)
headers2 = ["实验号", "日期", "内容", "配置", "关键结果", "判定"]
for c, h in enumerate(headers2, 2):
    ws2.cell(row=4, column=c, value=h)
B.style_header_row(ws2, 4, 2, 7)
runs = [
    ["E001", "2026-08-07~08", "LFBT 基线复现(冻结协议)", "seed 0, PTB-XL folds1-8 预训练", "LP 0.7177 / FT100 0.7173 / FT10 0.6431 AUPRC", "✅ 成立(偏差<1pt),基线零点"],
    ["E002", "2026-08-09~11", "AR-LFBT 候选模块筛选 B1-B4", "3 seed vs B0", "全部未稳定超 B0;B3(mean)持平留作对照", "❌ 关系层注意力/重加权路关闭"],
    ["E003", "2026-08-11~13", "六类跨领域模块单消融+自动复合", "GRN/Mixer/JEPA 等", "Mixer 明显退化;复合全败", "❌ 小模块堆叠不成,Mixer 否决 D1-GAT"],
    ["E004", "2026-08-13~17", "GRN-LFBT 专项验证", "5 seed+参数匹配", "FT-10 仅小趋势", "❌ 方法学终止(改 LRTC-Net 应用论文)"],
    ["E005", "2026-08-17~19", "RCLMD 跨导联掩码蒸馏 R01-R06", "完整师生框架", "R06 vs R01 clean -0.023", "❌ 蒸馏族证伪→D10 弃"],
    ["E006", "2026-08-24~09-08", "结构化生理-空间先验 P1/S1,2×4×3 seed", "24/24 全流程", "AR-B3 full LP +0.0040(3/3),AUPRC 0.7146 全场最高;B0 -0.0016(0/3)", "⚠️ 弱正未过 1pt 粗筛线,方向未证伪"],
    ["E007", "2026-09-08", "LRTC-Net(LGA) 2×2 正交开关 M0-M3", "3 seed(0,2,4) 配对", "M1-M0 ft100 -0.0089 CI[-0.0162,-0.0020] 显著负", "❌ 拉近式分组错误;轻量辅助项淹没于噪声"],
    ["E008", "2026-09-15 启动", "主线 S1: D9 VICReg 化粗筛(vicreg/bn/ln 三变体, seed 0/2)+batch 曲线 128/512,N4 附带", "PTB-XL 单库语料,B0 同配置", "官方系数 25/25/1: vicreg s0 0.6981(−1.96pt) / s2 0.6830(−3.47pt);+BN 0.6534(−6.43pt);+LN 0.6798(−3.79pt)——均 vs B0 0.7177;vicreg seed 方差 1.5pt >> B0 ±0.4pt", "粗筛全负;仅存下一步=系数重标定(D9-R1)+最小防塌补丁(D9-lite)+batch 曲线;B0 s2 配对基准 13:15"],
    ["N1", "2026-09-15 22:37 判定", "多库混合预训练(38,682条)生死判定", "seed 0, 200ep, batch 128, B0 原版目标", "LP AUPRC 0.6870 vs B0 0.7177 (Δ−3.07pt), AUROC −0.76pt", "❌ 不过线:不切换语料,G2 取消;混合比例/两阶段课程留作附录级改道"],
    ["W0-4", "2026-09-15", "服务器 LP 复现校验(本机 RTX4090, DL env torch2.5.1+cu121)", "checkpoint ptxl_gamma08, data/ptbxl 体系A, seed 0", "AUROC 0.9126 / AUPRC 0.7173 vs E001 锚点 0.7177,Δ=0.04pt", "✅ 通过(±0.5pt 内),管线与 E001 行为一致"],
]
for i, r in enumerate(runs):
    rn = 5 + i
    for c, v in enumerate(r, 2):
        ws2.cell(row=rn, column=c, value=v)
    B.style_data_row(ws2, rn, 2, 7, i)
B.auto_fit_columns(ws2, header_row=4, data_start_row=5)
B.auto_fit_row_heights(ws2, header_row=4, data_start_row=5, data_end_row=4 + len(runs))
ws2.freeze_panes = "C5"

# ---------- Sheet 3: 决策记录 ----------
ws3 = wb.create_sheet("决策记录")
B.setup_sheet(ws3, title="关键决策记录(04 号文档 §审计结论+01 号红线)", last_col=5)
headers3 = ["编号", "日期", "决策", "证据/理由"]
for c, h in enumerate(headers3, 2):
    ws3.cell(row=4, column=c, value=h)
B.style_header_row(ws3, 4, 2, 5)
dec = [
    ["#9", "2026-09-15", "下游划分定版:体系A(data/ptbxl + manifest)为冻结协议;体系B(data/downstream, data/pretrain)弃用", "随机划分 fold 混布、1513 患者跨 split、464 重复文件、177 条 train∩test 污染;runlog/W0/audit_data_report"],
    ["#NFH", "2026-09-15", "NFH/Chapman 私有弃用;SNPH 保密不可得", "语料扩展走 N1 本地多库;D8 鲁棒评估改道 SNR 曲线+37条件+跨库"],
    ["#算力", "2026-09-15", "正式训练协议按单卡 RTX 4090 24G、batch 128;本地仅开发/单测", "论文最优 batch 2048 不可达→D9 机会"],
    ["#W0-2", "2026-09-15", "服务器环境定版:conda DL env(Python3.10 + torch2.5.1+cu121 + wfdb/sklearn/pandas 齐全)", "nvidia-smi RTX4090 可见;其余候选 env 缺 torch 或版本不符"],
    ["#纪律", "—", "红线:唯一变量、开关全关==B0 逐位一致、两段式筛选(粗筛 1-2 seed Δ≥1.0pt)、5seed+配对检验+bootstrap CI、负对照/参数对齐自证、不选择性报告", "01 §六"],
]
for i, r in enumerate(dec):
    rn = 5 + i
    for c, v in enumerate(r, 2):
        ws3.cell(row=rn, column=c, value=v)
    B.style_data_row(ws3, rn, 2, 5, i)
B.auto_fit_columns(ws3, header_row=4, data_start_row=5)
B.auto_fit_row_heights(ws3, header_row=4, data_start_row=5, data_end_row=4 + len(dec))
ws3.freeze_panes = "C5"

wb.properties.creator = "Z.ai"
out = r"F:\新实验\LFBT实验方向清单.xlsx"
wb.save(out)
print("saved", out)
